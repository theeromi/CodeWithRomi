from io import BytesIO
from pathlib import Path

from pypdf import PdfReader
from docx import Document as DocxDocument


SUPPORTED_MIME = {
    "application/pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "text/plain",
    "text/markdown",
}


def is_supported(mime: str, filename: str) -> bool:
    if mime in SUPPORTED_MIME:
        return True
    ext = Path(filename).suffix.lower()
    return ext in {".pdf", ".docx", ".txt", ".md"}


def extract_text(path: str, mime: str, filename: str) -> str:
    ext = Path(filename).suffix.lower()
    if mime == "application/pdf" or ext == ".pdf":
        return _extract_pdf(path)
    if mime.endswith("wordprocessingml.document") or ext == ".docx":
        return _extract_docx(path)
    if mime.startswith("text/") or ext in {".txt", ".md"}:
        return Path(path).read_text(encoding="utf-8", errors="replace")
    raise ValueError(f"Unsupported file type: {mime} ({filename})")


def _extract_pdf(path: str) -> str:
    reader = PdfReader(path)
    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:
            pages.append("")
    return "\n\n".join(p.strip() for p in pages if p and p.strip())


def _extract_docx(path: str) -> str:
    doc = DocxDocument(path)
    parts = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text and cell.text.strip():
                    parts.append(cell.text.strip())
    return "\n".join(parts)
